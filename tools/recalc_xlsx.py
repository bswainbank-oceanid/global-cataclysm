"""
Recalculate an .xlsx with LibreOffice so its live formulas carry cached values
(openpyxl writes formulas but never evaluates them, so a workbook built by
tools/build_master_xlsx.py / tools/build_setup_tab.py shows blanks in any viewer
that doesn't recalculate on open). LibreOffice loads the file, computes every
formula and saves it back; formulas, formatting, conditional formats, data
validation and merged cells are preserved.

    python tools/recalc_xlsx.py [exports/GC1972_Territories.xlsx]

Finds soffice on PATH or in the usual install folders (Windows / macOS / Linux).
Replaces the file in place; leaves it untouched if LibreOffice fails or the result
has fewer formulas than the input.
"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else 'exports/GC1972_Territories.xlsx'

CANDIDATES = [
    shutil.which('soffice'), shutil.which('libreoffice'),
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
    '/Applications/LibreOffice.app/Contents/MacOS/soffice',
    '/usr/bin/soffice', '/usr/local/bin/soffice',
]
soffice = next((c for c in CANDIDATES if c and os.path.exists(c)), None)
if soffice is None:
    sys.exit('LibreOffice (soffice) not found')


def formula_count(p):
    wb = load_workbook(p)
    return sum(1 for ws in wb for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith('='))


before = formula_count(path)
with tempfile.TemporaryDirectory() as tmp:
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path],
                   check=True, capture_output=True, timeout=300)
    out = os.path.join(tmp, os.path.basename(path))
    after = formula_count(out)
    if after < before:
        sys.exit(f'recalculated file has {after} formulas, input had {before}; leaving {path} untouched')
    shutil.copyfile(out, path)
print(f'recalculated {path} via LibreOffice ({after} formulas)')
