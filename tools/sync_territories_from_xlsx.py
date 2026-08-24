"""
Read exports/GC1972_Territories.xlsx's 'All Territories' tab and write any
edits (Faction / Value / SC / Name) back into data/territories.json, the
canonical source. This is the missing half of the loop: build_master_xlsx.py
already generates the xlsx FROM the JSON; this script is what makes editing
the xlsx by hand (instead of the JSON directly) actually stick.

What it does NOT do:
  - Add or remove territories. The xlsx's 87 rows must match the 87 land
    ids already in data/territories.json 1:1 -- a row with an id not in
    the JSON, or a JSON id missing from the sheet, is a validation error,
    not a silent add/delete. New territories need map coordinates and an
    adjacency-graph position that this script has no way to invent.
  - Recompute anything derived (foreign-neighbor flags, the 'distant' flag,
    stacking caps, default sea zones). A faction reassignment here can
    change all of those. Re-run the pipeline after syncing:
        python3 tools/sync_territories_from_xlsx.py
        python3 tools/build_all.py
  - Reconcile data/scenarios/starting_setup_200ipc.json. If you reassign a
    territory's faction here, any existing scenario purchase entries for
    that territory under its OLD faction become stale -- the Initial Setup
    tab will flag them (red highlight, mismatched-faction rows), but
    tools/validate_setup.py currently assumes every purchase's territory
    still belongs to the faction that bought it and will raise a KeyError
    rather than a clean validation error if that's no longer true. Treat a
    faction reassignment as the start of a scenario edit, not a
    fire-and-forget spreadsheet tweak.

Validates before writing anything -- Faction must be blank or one of the 6
codes, Value must be within data/rules.json's territory_value_range, ids
must match 1:1. On any validation error, nothing is written. Use --dry-run
to see what would change without writing.

Run from the repo root:
    python3 tools/sync_territories_from_xlsx.py [--dry-run]
    [--xlsx PATH] [--territories PATH]
"""
import argparse
import json
import sys

from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument('--xlsx', default='exports/GC1972_Territories.xlsx')
parser.add_argument('--territories', default='data/territories.json')
parser.add_argument('--dry-run', action='store_true', help="report what would change; don't write")
args = parser.parse_args()

rules = json.load(open('data/rules.json'))['setup']
VALUE_MIN, VALUE_MAX = rules['territory_value_range']
VALID_FACTIONS = set(json.load(open('data/factions.json'))['factions'].keys())
SC_TARGET = rules['strategic_centers_per_faction']

territories = json.load(open(args.territories))
land_by_id = {sp['id']: sp for sp in territories['spaces'] if sp['type'] == 'land'}

wb = load_workbook(args.xlsx, data_only=False)
ws = wb['All Territories']

# ---- read every data row (row 4 until the 'Total' row) ----
rows = []
r = 4
while True:
    id_cell = ws.cell(row=r, column=1).value
    if id_cell is None or id_cell == 'Total':
        break
    rows.append({
        'row': r,
        'id': id_cell,
        'name': ws.cell(row=r, column=2).value,
        'faction': ws.cell(row=r, column=3).value,
        'value': ws.cell(row=r, column=4).value,
        'sc': ws.cell(row=r, column=5).value,
    })
    r += 1

# ---- validate ----
errors = []
seen_ids = set()
sheet_ids = set()
for entry in rows:
    r = entry['row']
    tid = entry['id']
    if not isinstance(tid, int):
        errors.append(f"row {r}: Territory ID {tid!r} is not a whole number")
        continue
    if tid in seen_ids:
        errors.append(f"row {r}: duplicate Territory ID {tid}")
    seen_ids.add(tid)
    sheet_ids.add(tid)

    if tid not in land_by_id:
        errors.append(f"row {r}: Territory ID {tid} does not match any territory in {args.territories} "
                       f"(rows can't add new territories -- see the script's docstring)")

    fac = (entry['faction'] or '').strip() or None
    if fac is not None and fac not in VALID_FACTIONS:
        errors.append(f"row {r} (id {tid}): invalid Faction {entry['faction']!r} "
                       f"(must be blank or one of {sorted(VALID_FACTIONS)})")

    val = entry['value']
    if not isinstance(val, int) or not (VALUE_MIN <= val <= VALUE_MAX):
        errors.append(f"row {r} (id {tid}): Value {val!r} must be a whole number "
                       f"from {VALUE_MIN} to {VALUE_MAX}")

    sc = entry['sc']
    if sc not in (None, '', 'Yes'):
        errors.append(f"row {r} (id {tid}): SC must be blank or \"Yes\", got {sc!r}")

missing_ids = set(land_by_id.keys()) - sheet_ids
for tid in sorted(missing_ids):
    errors.append(f"{args.territories} has territory id {tid} ({land_by_id[tid]['name']}) "
                   f"with no matching row in the sheet")

if errors:
    print(f'{len(errors)} validation error(s) -- nothing written:')
    for e in errors:
        print(' -', e)
    sys.exit(1)

# ---- diff ----
changes = []
for entry in rows:
    tid = entry['id']
    sp = land_by_id[tid]
    new_name = entry['name']
    new_faction = (entry['faction'] or '').strip() or None
    new_value = entry['value']
    new_sc = entry['sc'] == 'Yes'

    for field, old, new in (
        ('name', sp['name'], new_name),
        ('faction', sp['faction'], new_faction),
        ('value', sp['value'], new_value),
        ('strategic_center', sp['strategic_center'], new_sc),
    ):
        if old != new:
            changes.append((tid, sp['name'], field, old, new))
            sp[field] = new

if not changes:
    print('No changes -- the xlsx matches data/territories.json already.')
    sys.exit(0)

print(f'{len(changes)} change(s){" (dry run, not written)" if args.dry_run else ""}:')
for tid, name, field, old, new in changes:
    print(f'  [{tid}] {name}: {field} {old!r} -> {new!r}')

# ---- SC-count sanity warning (non-fatal) ----
sc_counts = {}
for sp in land_by_id.values():
    if sp['faction']:
        sc_counts.setdefault(sp['faction'], 0)
        if sp['strategic_center']:
            sc_counts[sp['faction']] += 1
for fac in VALID_FACTIONS:
    count = sc_counts.get(fac, 0)
    if count != SC_TARGET:
        print(f'  warning: {fac} now has {count} Strategic Center(s), expected {SC_TARGET}')

if args.dry_run:
    print('\nDry run -- data/territories.json not written.')
    sys.exit(0)

with open(args.territories, 'w') as f:
    json.dump(territories, f, indent=2)
print(f'\nwrote {args.territories}. Now run: python3 tools/build_all.py')
